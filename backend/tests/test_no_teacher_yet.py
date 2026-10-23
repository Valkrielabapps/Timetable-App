"""A subject nobody can teach yet is timetabled anyway, marked "No teacher yet".

Before, one subject without a qualified teacher (Assembly, a new elective) failed
the whole school's generation. Now its lessons take their place in the
section's week with no teacher, and the admin is told which ones they are.
"""
import io

import openpyxl

from app.core.database import SessionLocal
from app.models.school import ClassGroup, Period, Subject, Teacher, Timetable, TimetableEntry
from app.services.solver import generate_school_timetable
from tests.conftest import create_school, signup
from tests.test_ir_end_to_end import _class_group, _require, _school, _subject, _teacher, db  # noqa: F401


def test_lessons_with_no_teacher_are_still_placed(db):  # noqa: F811
    school = _school(db, days=1, per_day=4)
    eng = _subject(db, school, "English")
    assembly = _subject(db, school, "Assembly")
    rao = _teacher(db, school, "Rao", [eng.id])
    a = _class_group(db, school, name="A")
    _require(db, a, eng, 2)
    _require(db, a, assembly, 2)

    result = generate_school_timetable(db, school.id)
    assert result.status in ("optimal", "feasible"), result.errors
    by_subject = {}
    for x in result.assignments:
        by_subject.setdefault(x["subject_id"], []).append(x)
    assert len(by_subject[assembly.id]) == 2
    assert all(x["teacher_id"] is None for x in by_subject[assembly.id])
    assert all(x["teacher_id"] == rao.id for x in by_subject[eng.id])
    # Never two things in one of the section's periods.
    assert len({x["period_id"] for x in result.assignments}) == 4
    assert any("Assembly" in w and "No teacher yet" in w for w in result.warnings), result.warnings


def test_a_lesson_with_no_teacher_books_nobody(db):  # noqa: F811
    """Two sections' teacherless lessons can share a period; they don't use
    up anyone's time."""
    school = _school(db, days=1, per_day=2)
    assembly = _subject(db, school, "Assembly")
    eng = _subject(db, school, "English")
    _teacher(db, school, "Rao", [eng.id])
    a = _class_group(db, school, name="A")
    b = _class_group(db, school, name="B")
    _require(db, a, assembly, 2)
    _require(db, b, assembly, 1)
    _require(db, b, eng, 1)

    result = generate_school_timetable(db, school.id)
    assert result.status in ("optimal", "feasible"), result.errors
    assert sum(1 for x in result.assignments if x["teacher_id"] is None) == 3


def test_a_locked_lesson_with_no_teacher_keeps_its_slot(db):  # noqa: F811
    school = _school(db, days=1, per_day=4)
    assembly = _subject(db, school, "Assembly")
    a = _class_group(db, school, name="A")
    _require(db, a, assembly, 1)
    last = db.query(Period).filter(Period.school_id == school.id, Period.order == 4).one()
    tt = Timetable(school_id=school.id, status="draft", solver_status="optimal")
    db.add(tt)
    db.commit()
    db.add(TimetableEntry(timetable_id=tt.id, class_group_id=a.id, subject_id=assembly.id,
                          teacher_id=None, period_id=last.id, locked=True))
    db.commit()

    result = generate_school_timetable(db, school.id)
    assert result.status in ("optimal", "feasible"), result.errors
    assert [x["period_id"] for x in result.assignments] == [last.id]
    assert (a.id, assembly.id, None, last.id) in result.locked_keys


def _world(client):
    _, headers = signup(client)
    sid = create_school(client, headers)["id"]
    s = SessionLocal()
    try:
        periods = [Period(school_id=sid, day_of_week=0, order=o) for o in (1, 2)]
        assembly = Subject(school_id=sid, name="Assembly")
        eng = Subject(school_id=sid, name="English")
        rao = Teacher(school_id=sid, name="Rao", qualified_subject_ids=[])
        a = ClassGroup(school_id=sid, grade="Grade 8", name="A")
        b = ClassGroup(school_id=sid, grade="Grade 8", name="B")
        s.add_all([*periods, assembly, eng, rao, a, b])
        s.commit()
        tt = Timetable(school_id=sid, status="draft", solver_status="optimal")
        s.add(tt)
        s.commit()
        s.add_all([
            TimetableEntry(timetable_id=tt.id, class_group_id=a.id, subject_id=assembly.id,
                           teacher_id=None, period_id=periods[0].id),
            TimetableEntry(timetable_id=tt.id, class_group_id=a.id, subject_id=eng.id,
                           teacher_id=rao.id, period_id=periods[1].id),
            TimetableEntry(timetable_id=tt.id, class_group_id=b.id, subject_id=assembly.id,
                           teacher_id=None, period_id=periods[1].id),
        ])
        s.commit()
        return {"tt": tt.id, "p1": periods[0].id}, headers
    finally:
        s.close()


def test_the_timetable_and_exports_show_no_teacher_yet(client):
    w, h = _world(client)
    body = client.get(f"/api/timetables/{w['tt']}", headers=h)
    assert body.status_code == 200, body.text
    rows = [e for e in body.json()["entries"] if e["subject_name"] == "Assembly"]
    assert len(rows) == 2 and all(e["teacher_id"] is None and e["teacher_name"] is None for e in rows)

    xlsx = client.get(f"/api/timetables/{w['tt']}/export?format=xlsx", headers=h)
    assert xlsx.status_code == 200, xlsx.text
    wb = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    text = "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
    assert "No teacher yet" in text
    assert "Rao" in wb.sheetnames

    pdf = client.get(f"/api/timetables/{w['tt']}/export?format=pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"


def test_two_teacherless_lessons_can_share_a_period_when_moved(client):
    """Moving B's Assembly next to A's isn't "the teacher is already busy"."""
    w, h = _world(client)
    entries = client.get(f"/api/timetables/{w['tt']}", headers=h).json()["entries"]
    b_assembly = next(e for e in entries if e["subject_name"] == "Assembly" and e["period_id"] != w["p1"])
    r = client.patch(f"/api/timetables/entries/{b_assembly['id']}", headers=h, json={"period_id": w["p1"]})
    assert r.status_code == 200, r.text
