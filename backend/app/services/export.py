"""
Export a generated timetable to Excel (.xlsx) or PDF.

Both formats show the same two views already in the frontend
(TimetableTab.jsx): one grid per class-group section ("By Section") and
one grid per teacher ("By Teacher", their combined schedule across every
section they teach) — so what an admin sees on screen, and what they hand
out to a section or a teacher on paper, always match.

Shared data-gathering happens once in `_load_grid_data` so the Excel and
PDF builders can never disagree about what's actually in the timetable —
they just render the same grid data two different ways.
"""
import io
import re
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.models.school import ClassGroup, ElectiveBlock, Period, Room, Subject, Teacher, Timetable, ZeroPeriod
from app.services.timetable_structure import lesson_numbers

DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

HEADER_FILL = "1E293B"  # slate-800, matches the frontend's dark header rows


def _natural_key(text: str) -> list:
    """Sort key that orders "Grade 2" before "Grade 10" and is case-blind."""
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", text)]


def _load_grid_data(db: Session, timetable: Timetable) -> dict:
    school_id = timetable.school_id
    periods = db.query(Period).filter(Period.school_id == school_id).all()
    days = sorted({p.day_of_week for p in periods})
    orders = sorted({p.order for p in periods})
    period_by_day_order = {(p.day_of_week, p.order): p for p in periods}

    # Pages come out in the order a school reads them: grades in the sidebar's
    # order (the admin's custom grade order if set, else natural order so
    # "Grade 9" comes before "Grade 10"), sections A, B, C within each grade,
    # then teachers alphabetically. Unsorted, the database's own row order
    # gave 9-B, 10-B, 9-C...
    from app.models.school import School

    school = db.get(School, school_id)
    grade_rank = {g: i for i, g in enumerate((school.grade_order or []) if school else [])}
    class_groups = sorted(
        db.query(ClassGroup).filter(ClassGroup.school_id == school_id).all(),
        key=lambda cg: (
            cg.grade is None,  # ungrouped sections last
            grade_rank.get(cg.grade, len(grade_rank)),
            _natural_key(cg.grade or ""),
            _natural_key(cg.name or ""),
        ),
    )
    teachers = sorted(
        db.query(Teacher).filter(Teacher.school_id == school_id).all(),
        key=lambda t: _natural_key(t.name or ""),
    )
    subjects_by_id = {s.id: s for s in db.query(Subject).filter(Subject.school_id == school_id).all()}
    teachers_by_id = {t.id: t for t in teachers}
    class_groups_by_id = {c.id: c for c in class_groups}
    rooms_by_id = {r.id: r for r in db.query(Room).filter(Room.school_id == school_id).all()}
    block_names = {
        b.id: b.name for b in db.query(ElectiveBlock).filter(ElectiveBlock.school_id == school_id).all()
    }

    # Rows are labelled by lesson number, not row position: "Period 1" is the
    # first lesson even when a class teacher period comes before it.
    numbers = lesson_numbers(periods)
    zero_grades: dict[int, set] = defaultdict(set)
    for z in db.query(ZeroPeriod).filter(ZeroPeriod.school_id == school_id).all():
        zero_grades[z.period_id].add(z.grade)

    entries_by_class_period = defaultdict(list)
    entries_by_teacher_period = defaultdict(list)
    for e in timetable.entries:
        entries_by_class_period[(e.class_group_id, e.period_id)].append(e)
        entries_by_teacher_period[(e.teacher_id, e.period_id)].append(e)

    return {
        "days": days,
        "orders": orders,
        "period_by_day_order": period_by_day_order,
        "class_groups": class_groups,
        "teachers": teachers,
        "subjects_by_id": subjects_by_id,
        "teachers_by_id": teachers_by_id,
        "class_groups_by_id": class_groups_by_id,
        "rooms_by_id": rooms_by_id,
        "block_names": block_names,
        "numbers": numbers,
        "zero_grades": zero_grades,
        "entries_by_class_period": entries_by_class_period,
        "entries_by_teacher_period": entries_by_teacher_period,
    }


def _class_group_label(cg: ClassGroup) -> str:
    return f"{cg.grade} - {cg.name}" if cg.grade else cg.name


def _row_label(data: dict, order: int) -> str:
    """"Period 3" for a lesson row; the slot's own name for anything else."""
    present = [data["period_by_day_order"].get((d, order)) for d in data["days"]]
    present = [p for p in present if p]
    lesson = next((p for p in present if p.id in data["numbers"]), None)
    if lesson is not None:
        return f"Period {data['numbers'][lesson.id]}"
    kind = present[0].kind if present else "break"
    if kind == "break":
        return next((p.label for p in present if p.label), None) or "Break"
    return {"class_teacher": "Class teacher", "zero": "Zero period"}.get(kind, "Break")


def _special(data: dict, period, key, mode: str) -> list[tuple[str, str]] | None:
    """What a slot shows when it isn't an ordinary lesson: (title, who) pairs,
    or None for a lesson slot that has lessons as usual.

    The class teacher period shows each section with its class teacher - it is
    never in the timetable's entries, since nobody solves it. A zero period
    kept free for a section's grade shows as such, so the gap reads as
    intended rather than as a free period someone forgot to fill.
    """
    if period.kind == "class_teacher":
        if mode == "section":
            cg = data["class_groups_by_id"].get(key)
            teacher = data["teachers_by_id"].get(cg.class_teacher_id) if cg and cg.class_teacher_id else None
            return [("Class teacher", teacher.name if teacher else "")]
        sections = [
            _class_group_label(cg) for cg in data["class_groups"] if cg.class_teacher_id == key
        ]
        return [("Class teacher", ", ".join(sorted(sections)))] if sections else []
    if period.kind == "zero":
        return [("Zero period", "")]
    if period.kind == "break":
        return []
    if mode == "section":
        cg = data["class_groups_by_id"].get(key)
        grades = data["zero_grades"].get(period.id, set())
        if cg and (None in grades or cg.grade in grades):
            return [("Zero period", "")]
    return None


def _cell_text(entry, data: dict, mode: str) -> str:
    """mode='section': shows subject + teacher (for a section's own grid).
    mode='teacher': shows subject + section (for a teacher's own grid).
    Both also show the assigned room, if any (see solver.py's best-effort
    room assignment — not every entry is guaranteed a room)."""
    subject = data["subjects_by_id"].get(entry.subject_id)
    subject_name = subject.name if subject else "?"
    room = data["rooms_by_id"].get(entry.room_id) if entry.room_id else None
    room_suffix = f"\n[{room.name}]" if room else ""
    if mode == "section":
        teacher = data["teachers_by_id"].get(entry.teacher_id)
        return f"{subject_name}\n({teacher.name if teacher else '?'}){room_suffix}"
    cg = data["class_groups_by_id"].get(entry.class_group_id)
    return f"{subject_name}\n({_class_group_label(cg) if cg else '?'}){room_suffix}"


def _grid_rows(data: dict, entries_index: dict, entry_key_prefix, mode: str) -> list[list[str]]:
    """Builds a list of rows (including the header row) for one grid:
    row[0] = "Period N", followed by one cell per day. `entries_index` is
    either entries_by_class_period or entries_by_teacher_period;
    entry_key_prefix is the class_group_id or teacher_id to look up."""
    rows = [["Period"] + [DAY_NAMES[d] for d in data["days"]]]
    for order in data["orders"]:
        row = [_row_label(data, order)]
        for day in data["days"]:
            period = data["period_by_day_order"].get((day, order))
            text = ""
            special = _special(data, period, entry_key_prefix, mode) if period else None
            if special is not None:
                text = "\n".join(f"{title}\n({who})" if who else title for title, who in special)
            elif period:
                entries = entries_index.get((entry_key_prefix, period.id), [])
                text = "\n".join(_cell_text(e, data, mode) for e in entries)
                # A section's block slot holds every option at once; naming
                # the block says why one cell has three subjects in it.
                block_id = entries[0].elective_block_id if entries else None
                if mode == "section" and block_id is not None and len(entries) > 1:
                    text = f"{data['block_names'].get(block_id, 'Elective block')}\n{text}"
            row.append(text)
        rows.append(row)
    return rows


def _safe_sheet_name(name: str) -> str:
    """Excel sheet names can't contain : \\ / ? * [ ] and are capped at 31 chars."""
    cleaned = re.sub(r"[:\\/?*\[\]]", "-", name)
    return cleaned[:31] or "Sheet"


def _has_content(rows: list[list[str]]) -> bool:
    return any(any(r[1:]) for r in rows[1:])


def _school_grids(data: dict) -> list[tuple[str, str, list[list[str]]]]:
    """(sheet name, page title, rows) for every section, then every teacher."""
    grids = []
    for cg in data["class_groups"]:
        rows = _grid_rows(data, data["entries_by_class_period"], cg.id, mode="section")
        if _has_content(rows):
            label = _class_group_label(cg)
            grids.append((label, f"{label} — Timetable", rows))
    for teacher in data["teachers"]:
        rows = _grid_rows(data, data["entries_by_teacher_period"], teacher.id, mode="teacher")
        if _has_content(rows):
            grids.append((teacher.name, f"{teacher.name} — Teaching Schedule", rows))
    return grids


def _student_grids(data: dict, class_group: ClassGroup, combinations) -> list[tuple[str, str, list[list[str]]]]:
    """One grid per combination of elective choices, for one section.

    Each is the section's own timetable filtered to what that student
    attends: the common subjects, plus only the chosen option from each block.
    Numbered in the sheet name because Excel caps names at 31 characters, and
    "Physics + Chemistry + Mathematics" truncated would collide with its
    neighbours.
    """
    label = _class_group_label(class_group)
    grids = []
    for number, combination in enumerate(combinations, start=1):
        index: dict = defaultdict(list)
        for (cg_id, period_id), entries in data["entries_by_class_period"].items():
            if cg_id == class_group.id:
                index[(cg_id, period_id)] = [e for e in entries if combination.keeps(e)]
        rows = _grid_rows(data, index, class_group.id, mode="section")
        grids.append((
            f"{number}. {combination.label}",
            f"{label} — {combination.label}",
            rows,
        ))
    return grids


def _excel(grids, school_name: str = "", free_plan: bool = False) -> bytes:
    """Same design as the PDF (see _pdf below): heading block, black header
    row, shaded period column, striped rows of equal height, each lesson as
    a bold subject over a muted teacher/section and room, breaks as one
    green band across the week. Set up to print on one landscape page.

    Cells use rich text for the bold/muted split. A reader that ignores rich
    text (or openpyxl loading the file back) still sees the same plain text,
    one item per line."""
    from datetime import date

    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Border, Side
    from openpyxl.worksheet.properties import PageSetupProperties

    generated = date.today().strftime("%d %b %Y").lstrip("0")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    thin = Side(style="thin", color="E5E5E5")
    frame = Side(style="thin", color="D4D4D4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = lambda hex_: PatternFill("solid", fgColor=hex_)  # noqa: E731
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_mid = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

    subj_font = InlineFont(rFont="Calibri", b=True, sz=11, color="171717")
    who_font = InlineFont(rFont="Calibri", sz=9, color="525252")
    room_font = InlineFont(rFont="Calibri", sz=8, color="737373")
    block_font = InlineFont(rFont="Calibri", b=True, sz=8, color="047857")

    for sheet_name, _title, rows, meta in grids:
        ws = wb.create_sheet(_safe_sheet_name(sheet_name))
        ws.sheet_view.showGridLines = False
        n_days = len(rows[0]) - 1
        last_col = n_days + 1

        # Heading block
        ws.cell(row=1, column=1, value=meta["heading"]).font = Font(name="Calibri", bold=True, size=18, color="171717")
        ws.cell(row=2, column=1, value=meta["kind"]).font = Font(name="Calibri", size=10, color="737373")
        brand = ws.cell(row=1, column=last_col, value="Timetablz")
        brand.font = Font(name="Calibri", bold=True, size=10, color="171717")
        brand.alignment = Alignment(horizontal="right", vertical="center")
        if school_name:
            sc = ws.cell(row=2, column=last_col, value=school_name)
            sc.font = Font(name="Calibri", size=10, color="737373")
            sc.alignment = Alignment(horizontal="right")
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[3].height = 8

        # Header row
        head_row = 4
        for col, label in enumerate(rows[0], start=1):
            c = ws.cell(row=head_row, column=col, value=label)
            c.font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
            c.fill = fill("171717")
            c.alignment = centre
            c.border = Border(left=Side(style="thin", color="404040"), right=Side(style="thin", color="404040"))
        ws.row_dimensions[head_row].height = 26
        ws.column_dimensions["A"].width = 14
        for col in range(2, last_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = 24

        teaching_index = 0
        r = head_row
        for row in meta["cells"]:
            r += 1
            if row["is_break"]:
                ws.cell(row=r, column=1).fill = fill("ECFDF5")
                ws.cell(row=r, column=1).border = border
                band = ws.cell(row=r, column=2, value=row["label"])
                band.font = Font(name="Calibri", italic=True, size=10, color="047857")
                band.alignment = centre
                for col in range(2, last_col + 1):
                    ws.cell(row=r, column=col).fill = fill("ECFDF5")
                    ws.cell(row=r, column=col).border = border
                if n_days > 1:
                    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last_col)
                ws.row_dimensions[r].height = 20
                continue

            stripe = "FAFAFA" if teaching_index % 2 else "FFFFFF"
            teaching_index += 1
            p = ws.cell(row=r, column=1, value=row["label"])
            p.font = Font(name="Calibri", bold=True, size=10, color="171717")
            p.fill = fill("F5F5F5")
            p.alignment = left_mid
            p.border = border

            lines_needed = 2
            for d, cell in enumerate(row["cells"]):
                c = ws.cell(row=r, column=d + 2)
                c.fill = fill(stripe)
                c.alignment = centre
                c.border = border
                blocks = []
                if cell["block"]:
                    blocks.append(TextBlock(block_font, cell["block"]))
                for subject, who, room in cell["lessons"]:
                    if blocks:
                        blocks.append(TextBlock(who_font, "\n"))
                    blocks.append(TextBlock(subj_font, subject))
                    if who:
                        blocks.append(TextBlock(who_font, "\n" + who))
                    if room:
                        blocks.append(TextBlock(room_font, "\n" + room))
                if blocks:
                    c.value = CellRichText(*blocks)
                    lines = (1 if cell["block"] else 0) + sum(1 + bool(w) + bool(rm) for _s, w, rm in cell["lessons"])
                    lines_needed = max(lines_needed, lines)
                elif cell["exists"] and meta["mode"] == "teacher":
                    c.value = "Free"
                    c.font = Font(name="Calibri", size=9, color="A3A3A3")
            ws.row_dimensions[r].height = max(40, lines_needed * 14 + 10)

        # Outer frame round the whole table
        for col in range(1, last_col + 1):
            top = ws.cell(row=head_row, column=col)
            bottom = ws.cell(row=r, column=col)
            bottom.border = Border(left=bottom.border.left, right=bottom.border.right, top=bottom.border.top, bottom=frame)
            top.border = Border(left=top.border.left, right=top.border.right, top=frame, bottom=top.border.bottom)

        ws.freeze_panes = ws.cell(row=head_row + 1, column=2)

        # Print: one landscape page, centred, with the same footer as the PDF
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.5
        ws.page_margins.top = ws.page_margins.bottom = 0.6
        ws.print_area = f"A1:{get_column_letter(last_col)}{r}"
        ws.oddFooter.left.text = (
            f"Made with Timetablz Free  ·  timetablz.com  ·  {generated}" if free_plan
            else f"Generated {generated}  ·  timetablz.com"
        )
        ws.oddFooter.left.size = 8
        ws.oddFooter.left.color = "A3A3A3"
        ws.oddFooter.right.text = "Page &P"
        ws.oddFooter.right.size = 8
        ws.oddFooter.right.color = "A3A3A3"

    if not wb.sheetnames:
        wb.create_sheet("Timetable").append(["This timetable has no entries yet."])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
#
# One page per grid, laid out like a printed school timetable rather than a
# spreadsheet dump: a heading block, a dark header row, a shaded period
# column, every lesson as subject (bold) over teacher/section and room
# (muted), centred in rows of equal height that fill the page, and breaks as
# a single band across the week. Colours are the app's own neutral palette
# with the brand green used only for breaks and elective-block labels.

_INK = colors.HexColor("#171717")       # neutral-900, the app's primary dark
_BODY = colors.HexColor("#404040")      # neutral-700
_MUTED = colors.HexColor("#737373")     # neutral-500
_FAINT = colors.HexColor("#A3A3A3")     # neutral-400
_LINE = colors.HexColor("#E5E5E5")      # neutral-200
_FRAME = colors.HexColor("#D4D4D4")     # neutral-300
_STRIPE = colors.HexColor("#FAFAFA")    # neutral-50
_PERIOD_COL = colors.HexColor("#F5F5F5")  # neutral-100
_GREEN = colors.HexColor("#047857")     # emerald-700, readable on white
_GREEN_TINT = colors.HexColor("#ECFDF5")  # emerald-50

_PAGE = landscape(A4)
_MARGIN_X = 1.4 * cm
_MARGIN_TOP = 1.9 * cm
_MARGIN_BOTTOM = 1.5 * cm


def _esc(text: str) -> str:
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _pdf_cells(data: dict, entries_index: dict, key, mode: str) -> list[dict]:
    """The same grid as _grid_rows, kept structured for the PDF: one dict per
    period order with its label, whether it's a break, and per day the list
    of (subject, who, room) lessons plus the elective block name if any."""
    rows = []
    for order in data["orders"]:
        day_periods = [data["period_by_day_order"].get((d, order)) for d in data["days"]]
        present = [p for p in day_periods if p]
        # A band across the whole row: a break, or a zero period for the whole
        # school. The class teacher period is not one - it names a different
        # teacher in each section, so it keeps its cells.
        is_break = bool(present) and all(p.kind in ("break", "zero") for p in present)
        cells = []
        for period in day_periods:
            lessons, block = [], None
            special = _special(data, period, key, mode) if period else None
            if special is not None:
                lessons = [(title, who, "") for title, who in special]
            elif period and not period.is_break:
                entries = entries_index.get((key, period.id), [])
                for e in entries:
                    subject = data["subjects_by_id"].get(e.subject_id)
                    if mode == "section":
                        who = data["teachers_by_id"].get(e.teacher_id)
                        who = who.name if who else ""
                    else:
                        cg = data["class_groups_by_id"].get(e.class_group_id)
                        who = _class_group_label(cg) if cg else ""
                    room = data["rooms_by_id"].get(e.room_id) if e.room_id else None
                    lessons.append((subject.name if subject else "?", who, room.name if room else ""))
                if mode == "section" and entries and entries[0].elective_block_id is not None and len(entries) > 1:
                    block = data["block_names"].get(entries[0].elective_block_id, "Elective block")
            cells.append({"lessons": lessons, "block": block, "exists": period is not None})
        rows.append({
            "label": _row_label(data, order),
            "is_break": is_break,
            "cells": cells,
        })
    return rows


def _pdf(grids, n_days: int, school_name: str = "", free_plan: bool = False) -> bytes:
    from datetime import date

    buf = io.BytesIO()
    generated = date.today().strftime("%d %b %Y").lstrip("0")

    def decorate(canvas, doc):
        width, height = _PAGE
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 10)
        canvas.setFillColor(_INK)
        canvas.drawString(_MARGIN_X, height - 1.05 * cm, "Timetablz")
        canvas.setFont("Helvetica", 9)
        canvas.setFillColor(_MUTED)
        if school_name:
            canvas.drawRightString(width - _MARGIN_X, height - 1.05 * cm, school_name)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(_FAINT)
        canvas.drawString(_MARGIN_X, 0.8 * cm, f"Generated {generated}  ·  timetablz.com")
        if free_plan:
            # Free-plan exports carry the product name, centred in the footer.
            canvas.setFont("Helvetica-Bold", 8)
            canvas.setFillColor(_MUTED)
            canvas.drawCentredString(width / 2, 0.8 * cm, "Made with Timetablz Free")
        canvas.drawRightString(width - _MARGIN_X, 0.8 * cm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf, pagesize=_PAGE,
        leftMargin=_MARGIN_X, rightMargin=_MARGIN_X, topMargin=_MARGIN_TOP, bottomMargin=_MARGIN_BOTTOM,
        title=school_name or "Timetable", author="Timetablz",
    )
    heading = ParagraphStyle("H", fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=_INK)
    sub = ParagraphStyle("S", fontName="Helvetica", fontSize=9.5, leading=12, textColor=_MUTED, spaceBefore=2, spaceAfter=12)
    head_cell = ParagraphStyle("HC", fontName="Helvetica-Bold", fontSize=9.5, leading=11, textColor=colors.white, alignment=1)
    period_cell = ParagraphStyle("PC", fontName="Helvetica-Bold", fontSize=9, leading=11, textColor=_INK)
    lesson = ParagraphStyle("L", fontName="Helvetica", fontSize=8, leading=10.5, textColor=_BODY, alignment=1)
    free = ParagraphStyle("F", parent=lesson, textColor=_FAINT)
    break_style = ParagraphStyle("B", fontName="Helvetica-Oblique", fontSize=8.5, leading=10, textColor=_GREEN, alignment=1)

    usable_w = _PAGE[0] - 2 * _MARGIN_X - 12  # frame padding, so the table lines up with the heading
    period_w = 2.6 * cm
    day_w = (usable_w - period_w) / max(n_days, 1)
    # Frame height, less the frame's own 6pt padding top and bottom, the
    # heading block (24pt title + 2 + 12 + 12pt subtitle) and a little slack,
    # so a normal week always fits on one page.
    usable_h = _PAGE[1] - _MARGIN_TOP - _MARGIN_BOTTOM - 12 - 50 - 8
    header_h = 0.85 * cm

    story = []
    for i, grid in enumerate(grids):
        _sheet, _title, rows, meta = grid
        if i:
            story.append(PageBreak())
        story.append(Paragraph(_esc(meta["heading"]), heading))
        story.append(Paragraph(_esc(meta["kind"]), sub))

        cells_rows = meta["cells"]
        days_header = rows[0][1:]
        table_rows = [[Paragraph("Period", head_cell)] + [Paragraph(_esc(d), head_cell) for d in days_header]]
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), _INK),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 1), (0, -1), 10),
            ("BOX", (0, 0), (-1, -1), 0.75, _FRAME),
            ("INNERGRID", (0, 1), (-1, -1), 0.5, _LINE),
            ("LINEAFTER", (0, 0), (-2, 0), 0.5, colors.HexColor("#404040")),
        ]

        n_breaks = sum(1 for r in cells_rows if r["is_break"])
        n_lessons = max(len(cells_rows) - n_breaks, 1)
        break_h = 0.6 * cm
        base_h = min(1.6 * cm, max(1.0 * cm, (usable_h - header_h - n_breaks * break_h) / n_lessons))
        heights = [header_h]
        teaching_index = 0

        for r_i, row in enumerate(cells_rows, start=1):
            if row["is_break"]:
                table_rows.append(["", Paragraph(_esc(row["label"]), break_style)] + [""] * (n_days - 1))
                style += [
                    ("SPAN", (1, r_i), (-1, r_i)),
                    ("BACKGROUND", (0, r_i), (-1, r_i), _GREEN_TINT),
                ]
                heights.append(break_h)
                continue

            line_count = 2
            out = [Paragraph(_esc(row["label"]), period_cell)]
            for cell in row["cells"]:
                parts = []
                if cell["block"]:
                    parts.append(f'<font color="#047857" size="7"><b>{_esc(cell["block"].upper())}</b></font>')
                for subject, who, room in cell["lessons"]:
                    piece = f'<font color="#171717" size="9"><b>{_esc(subject)}</b></font>'
                    if who:
                        piece += f"<br/>{_esc(who)}"
                    if room:
                        piece += f'<br/><font color="#737373" size="7.5">{_esc(room)}</font>'
                    parts.append(piece)
                if parts:
                    out.append(Paragraph("<br/>".join(parts), lesson))
                    lines = (1 if cell["block"] else 0) + sum(1 + bool(w) + bool(r) for _s, w, r in cell["lessons"])
                    line_count = max(line_count, lines)
                elif cell["exists"] and meta["mode"] == "teacher":
                    out.append(Paragraph("Free", free))
                else:
                    out.append("")
            table_rows.append(out)
            stripe = _STRIPE if teaching_index % 2 else colors.white
            style += [
                ("BACKGROUND", (1, r_i), (-1, r_i), stripe),
                ("BACKGROUND", (0, r_i), (0, r_i), _PERIOD_COL),
            ]
            teaching_index += 1
            heights.append(max(base_h, line_count * 10.5 + 12))

        table = Table(table_rows, colWidths=[period_w] + [day_w] * n_days, rowHeights=heights, repeatRows=1)
        table.setStyle(TableStyle(style))
        story.append(table)

    if not story:
        story.append(Paragraph("This timetable has no entries yet.", sub))

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()


def _with_pdf_meta(data: dict, grids, kinds) -> list:
    """Attach the structured PDF cells to each (sheet, title, rows) grid."""
    out = []
    for (sheet, title, rows), (heading, kind, index, key, mode) in zip(grids, kinds):
        out.append((sheet, title, rows, {
            "heading": heading, "kind": kind, "mode": mode,
            "cells": _pdf_cells(data, index, key, mode),
        }))
    return out


def _school_name(db: Session, timetable: Timetable) -> str:
    from app.models.school import School
    school = db.get(School, timetable.school_id)
    return school.name if school else ""


def _free_plan(db: Session, timetable: Timetable) -> bool:
    from app.models.school import School
    from app.services import plans
    school = db.get(School, timetable.school_id)
    return school is not None and not plans.clean_exports(school)


def _school_grids_with_meta(data: dict) -> list:
    kinds = []
    for cg in data["class_groups"]:
        if _has_content(_grid_rows(data, data["entries_by_class_period"], cg.id, mode="section")):
            kinds.append((_class_group_label(cg), "Class timetable", data["entries_by_class_period"], cg.id, "section"))
    for t in data["teachers"]:
        if _has_content(_grid_rows(data, data["entries_by_teacher_period"], t.id, mode="teacher")):
            kinds.append((t.name, "Teaching schedule", data["entries_by_teacher_period"], t.id, "teacher"))
    return _with_pdf_meta(data, _school_grids(data), kinds)


def _student_grids_with_meta(data: dict, class_group: ClassGroup, combinations) -> list:
    label = _class_group_label(class_group)
    kinds = []
    for combination in combinations:
        index: dict = defaultdict(list)
        for (cg_id, period_id), entries in data["entries_by_class_period"].items():
            if cg_id == class_group.id:
                index[(cg_id, period_id)] = [e for e in entries if combination.keeps(e)]
        kinds.append((label, f"Student timetable  ·  {combination.label}", index, class_group.id, "section"))
    return _with_pdf_meta(data, _student_grids(data, class_group, combinations), kinds)


def build_excel(db: Session, timetable: Timetable) -> bytes:
    data = _load_grid_data(db, timetable)
    return _excel(_school_grids_with_meta(data), _school_name(db, timetable), _free_plan(db, timetable))


def build_pdf(db: Session, timetable: Timetable) -> bytes:
    data = _load_grid_data(db, timetable)
    return _pdf(_school_grids_with_meta(data), len(data["days"]), _school_name(db, timetable), _free_plan(db, timetable))


def build_student_excel(db: Session, timetable: Timetable, class_group: ClassGroup, combinations) -> bytes:
    """One sheet per combination of elective choices in one section."""
    data = _load_grid_data(db, timetable)
    return _excel(_student_grids_with_meta(data, class_group, combinations), _school_name(db, timetable), _free_plan(db, timetable))


def build_student_pdf(db: Session, timetable: Timetable, class_group: ClassGroup, combinations) -> bytes:
    """One page per combination of elective choices in one section."""
    data = _load_grid_data(db, timetable)
    return _pdf(_student_grids_with_meta(data, class_group, combinations), len(data["days"]), _school_name(db, timetable), _free_plan(db, timetable))
